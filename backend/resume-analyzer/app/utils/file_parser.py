import io
import logging
from pathlib import Path
from typing import Tuple

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text.strip())
        return "\n\n".join(text_parts)
    except ImportError:
        raise RuntimeError("pypdf not installed. Run: pip install pypdf")
    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        raise ValueError(f"Could not extract text from PDF: {e}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise ValueError(f"Could not extract text from DOCX: {e}")


def extract_text_from_file(
    file_bytes: bytes, filename: str
) -> Tuple[str, str]:
    """
    Extract text from uploaded file.
    Returns (extracted_text, detected_format).
    """
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(file_bytes), "pdf"
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(file_bytes), "docx"
    elif suffix == ".txt":
        try:
            return file_bytes.decode("utf-8"), "txt"
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1"), "txt"
    else:
        raise ValueError(f"Unsupported file format: {suffix}")


def validate_file_size(file_bytes: bytes, max_mb: int = 5) -> None:
    """Raise ValueError if file exceeds size limit."""
    size_mb = len(file_bytes) / (1024 * 1024)
    if size_mb > max_mb:
        raise ValueError(
            f"File size ({size_mb:.1f} MB) exceeds limit ({max_mb} MB)"
        )
